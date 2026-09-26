#!/usr/bin/env python3
"""Regenerate every derived file of the EV-pack EOL cell, byte-for-byte reproducibly.

Inputs (hand-authored): evals/answer_key/*, plant_data/ev_pack_eol/{io_list,parameters,cause_effect}.csv,
scenarios/ce_bindings.yaml, st30/program.il.
Outputs: st10/st20 legacy.il (golden + seeded defects, documentation stripped), device_comments.csv,
gxw3/*.csv (UTF-16LE, GX Works3 listed-instruction / device-comment format), scenarios/ce_*.yaml,
io_list.xlsx, cause_effect.xlsx.

    python scripts/build_data.py             # write in place
    python scripts/build_data.py --out DIR   # write the same tree under DIR (reproducibility checks)
"""
from __future__ import annotations

import argparse
import csv
import io
import sys
import zipfile
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ladder_harness.cell import load_cell, read_comments_csv, write_comments_csv  # noqa: E402
from ladder_harness.melsec.gxw3_csv import comments_to_csv_bytes, program_to_csv_bytes  # noqa: E402
from ladder_harness.melsec.program import format_il, parse_il  # noqa: E402
from ladder_harness.scenarios.from_ce import generate  # noqa: E402
from ladder_harness.variants import make_variant, used_devices  # noqa: E402

CELL = Path("plant_data/ev_pack_eol")
KEY = ROOT / "evals" / "answer_key"
FIXED_TIME = (2026, 9, 26, 0, 0, 0)
LEGACY_DEFECTS = {"ST20": ["D1", "D3", "D4", "D5"], "ST10": ["F1"]}


def normalize_zip(path: Path) -> None:
    """Rewrite an xlsx with fixed entry timestamps (and openpyxl's save-time `modified`) for reproducible bytes."""
    import re
    with zipfile.ZipFile(path) as z:
        entries = [(info.filename, z.read(info.filename)) for info in z.infolist()]
    stamp = "%04d-%02d-%02dT%02d:%02d:%02dZ" % FIXED_TIME
    entries = [(name, re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                             rb"\g<1>" + stamp.encode() + rb"\g<2>", data) if name == "docProps/core.xml" else data)
               for name, data in entries]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=FIXED_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            z.writestr(info, data)
    path.write_bytes(buf.getvalue())


def write_xlsx(out: Path, io_rows: list[dict], params: list[dict], ce_rows: list[dict]) -> None:
    import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    stamp = datetime.datetime(*FIXED_TIME)
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="E8EAED")

    def sheet(ws, rows, cols, widths):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"

    wb = Workbook()
    wb.properties.created = wb.properties.modified = stamp
    wb.properties.creator = wb.properties.lastModifiedBy = "ladder-sdlc-harness build_data.py"
    ws = wb.active
    ws.title = "IO List"
    sheet(ws, io_rows, ["tag", "device", "description", "signal", "range", "units", "station", "safety", "notes"],
          [12, 8, 52, 12, 8, 12, 8, 9, 44])
    sheet(wb.create_sheet("Parameters"), params, ["parameter", "value", "units", "device", "station", "note"],
          [20, 8, 22, 26, 8, 44])
    wb.save(out / "io_list.xlsx")
    normalize_zip(out / "io_list.xlsx")

    wb = Workbook()
    wb.properties.created = wb.properties.modified = stamp
    wb.properties.creator = wb.properties.lastModifiedBy = "ladder-sdlc-harness build_data.py"
    ws = wb.active
    ws.title = "Matrix"
    effects = sorted({r["effect_tag"] for r in ce_rows})
    ws.append(["C&E id", "Station", "Cause tag", "Cause", *effects, "Response (ms)", "Safety"])
    for c in ws[1]:
        c.font, c.fill = bold, head_fill
        c.alignment = Alignment(text_rotation=90 if c.column > 4 and c.value in effects else 0, wrap_text=True)
    for r in ce_rows:
        ws.append([r["id"], r["station"], r["cause_tag"], r["cause"],
                   *[r["action"] if r["effect_tag"] == e else "" for e in effects],
                   r["response_ms"], r["safety_class"]])
    ws.column_dimensions["D"].width = 52
    ws.freeze_panes = "E2"
    sheet(wb.create_sheet("List"), ce_rows, list(ce_rows[0].keys()), [12, 7, 16, 48, 18, 48, 16, 12, 16, 16])
    wb.save(out / "cause_effect.xlsx")
    normalize_zip(out / "cause_effect.xlsx")


def read_csv(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def build(out_root: Path) -> list[Path]:
    src = ROOT / CELL
    out = out_root / CELL
    for sub in ("st10", "st20", "st30", "gxw3", "scenarios"):
        (out / sub).mkdir(parents=True, exist_ok=True)
    cell = load_cell(src, KEY)
    golden_comments = read_comments_csv(KEY / "golden_comments.csv")
    written: list[Path] = []

    for station, ids in LEGACY_DEFECTS.items():
        legacy, comments = make_variant(cell, station, ids)
        (out / station.lower() / "legacy.il").write_text(format_il(legacy), encoding="utf-8")
        write_comments_csv(out / station.lower() / "device_comments.csv", comments)
        (out / "gxw3" / f"{station}_legacy.csv").write_bytes(
            program_to_csv_bytes(legacy, project=f"{station}_legacy", module_type="FX5U"))
        (out / "gxw3" / f"{station}_comments.csv").write_bytes(comments_to_csv_bytes(comments))
        written += [out / station.lower() / "legacy.il", out / station.lower() / "device_comments.csv"]

    st30 = parse_il((src / "st30" / "program.il").read_text(encoding="utf-8"), "ST30")
    st30_comments = {d: golden_comments[d] for d in used_devices(st30) if d in golden_comments}
    write_comments_csv(out / "st30" / "device_comments.csv", st30_comments)
    if out != src:
        (out / "st30" / "program.il").write_text(format_il(st30), encoding="utf-8")
    (out / "gxw3" / "ST30.csv").write_bytes(program_to_csv_bytes(st30, project="ST30", module_type="FX5U"))
    (out / "gxw3" / "ST30_comments.csv").write_bytes(comments_to_csv_bytes(st30_comments))

    ce_rows = read_csv(src / "cause_effect.csv")
    bindings = yaml.safe_load((src / "scenarios" / "ce_bindings.yaml").read_text(encoding="utf-8"))
    for station, suite in generate(ce_rows, bindings).items():
        path = out / "scenarios" / f"ce_{station.lower()}.yaml"
        text = ("# GENERATED by scripts/build_data.py from cause_effect.csv + scenarios/ce_bindings.yaml.\n"
                "# Do not edit; change the C&E matrix or the bindings and rebuild.\n")
        path.write_text(text + yaml.safe_dump(suite, sort_keys=False, allow_unicode=True, width=110), encoding="utf-8")
        written.append(path)

    write_xlsx(out, read_csv(src / "io_list.csv"), read_csv(src / "parameters.csv"), ce_rows)
    return written


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=ROOT, help="root to write under (default: the repo)")
    args = ap.parse_args()
    for path in build(args.out.resolve()):
        print(path.relative_to(args.out.resolve()))


if __name__ == "__main__":
    main()
